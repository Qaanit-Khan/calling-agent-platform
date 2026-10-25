import io
import uuid
from datetime import UTC, datetime
import logging
from typing import Any, Dict, List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, update, delete
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
import openpyxl.styles

from app.database import get_db
from app.dependencies import verify_app_client
from app.models.agent import Agent
from app.models.questionnaire import Questionnaire

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/questionnaires", tags=["questionnaires"], dependencies=[Depends(verify_app_client)])


def map_q_to_doc(q: Questionnaire) -> Dict[str, Any]:
    return {
        "id": q.id,
        "questionnaire_id": q.id,
        "tenant_id": q.tenant_id,
        "agent_id": q.agent_id,
        "version_id": q.version_id or "v1",
        "title": q.title,
        "description": q.description,
        "questions": q.questions or [],
        "status": q.status or "draft",
        "is_published": q.is_published or False,
        "version_history": q.version_history or [],
        "created_at": q.created_at.isoformat() if q.created_at else datetime.now(UTC).isoformat(),
        "updated_at": q.updated_at.isoformat() if q.updated_at else datetime.now(UTC).isoformat(),
    }


class QuestionnairePayload(BaseModel):
    title: Optional[str] = "Questionnaire"
    description: Optional[str] = None
    agent_id: Optional[str] = None
    tenant_id: Optional[str] = None
    version_id: Optional[str] = "v1"
    questions: Optional[List[Dict[str, Any]]] = Field(default_factory=list)
    sections: Optional[List[Dict[str, Any]]] = None
    status: Optional[str] = "draft"
    is_published: Optional[bool] = False


def extract_questions_from_payload(payload: QuestionnairePayload) -> List[Dict[str, Any]]:
    if payload.questions:
        return payload.questions
    if payload.sections:
        extracted = []
        for sec in payload.sections:
            for q_item in sec.get("questions", []):
                extracted.append({
                    "question_id": q_item.get("id") or str(uuid.uuid4())[:8],
                    "text": q_item.get("text") or q_item.get("question") or q_item.get("title", ""),
                    "expected_answer": q_item.get("expected_answer", ""),
                    "type": q_item.get("type", "text"),
                    "options": q_item.get("options", []),
                })
        return extracted
    return []


@router.get("")
async def list_questionnaires(
    agent_id: Optional[str] = None,
    tenant_id: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Questionnaire).order_by(Questionnaire.created_at.desc())
    if agent_id:
        stmt = stmt.where(Questionnaire.agent_id == agent_id)
    if tenant_id:
        stmt = stmt.where(Questionnaire.tenant_id == tenant_id)
    res = await db.execute(stmt)
    items = res.scalars().all()
    return [map_q_to_doc(q) for q in items]


@router.get("/excel-template")
async def download_excel_template():
    """
    Generate and download sample Excel questionnaire template with Question and Expected Answer columns.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questionnaire"

    # Header styling
    headers = ["Question", "Expected Answer"]
    ws.append(headers)

    header_fill = openpyxl.styles.PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = openpyxl.styles.Font(color="FFFFFF", bold=True, size=11)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font

    # Sample reference rows
    sample_rows = [
        ["नमस्कार साहेब, आपण कोणकोणत्या पिकांचे नियोजन केले आहे का?", "गहू / हरभरा / चणा / मका"],
        ["साहेब, आपण गहू साधारण किती एकरमध्ये घेण्याचे नियोजन केले आहे?", "Any Answer / Number of acres"],
        ["त्या गव्हासाठी बियाणे खरेदी करण्याचे प्लॅनिंग किती दिवसांत करत आहात?", "Any Answer"],
        ["Samruddhi Agritech च्या गहू व्हरायटी बद्दल अधिक माहिती पाहिजे आहे का?", "हो / नाही"],
    ]
    for row in sample_rows:
        ws.append(row)

    ws.column_dimensions["A"].width = 65
    ws.column_dimensions["B"].width = 35

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    response_headers = {
        "Content-Disposition": "attachment; filename=Questionnaire_Template.xlsx",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    return StreamingResponse(
        stream,
        headers=response_headers,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.post("/import-excel")
async def import_excel_questionnaire(
    file: UploadFile = File(...),
    title: Optional[str] = Form(None),
    agent_id: Optional[str] = Form(None),
    tenant_id: Optional[str] = Form(None),
    save_to_db: Optional[bool] = Form(False),
    db: AsyncSession = Depends(get_db),
):
    """
    Validate and parse an uploaded Excel (.xlsx) file containing Question and Expected Answer columns.
    Returns parsed questions and validation warnings for preview.
    If save_to_db is True, creates a new Questionnaire record directly in PostgreSQL.
    """
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload a valid Microsoft Excel (.xlsx) file.",
        )

    contents = await file.read()
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large. Maximum supported file size is 10MB.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        logger.error(f"Failed to parse Excel file: {e}")
        raise HTTPException(
            status_code=400,
            detail=f"Invalid or corrupted Excel file: {str(e)}",
        )

    ws = wb.active
    if not ws:
        raise HTTPException(status_code=400, detail="Excel workbook contains no active sheet.")

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="The uploaded Excel sheet is empty.")

    header_row_idx = 0
    question_col_idx = 0
    expected_ans_col_idx = 1

    first_row = [str(c or "").strip().lower() for c in rows[0]]
    has_header = False

    for idx, col_val in enumerate(first_row):
        if any(kw in col_val for kw in ["question", "prompt", "प्रश्न", "सवाल"]):
            question_col_idx = idx
            has_header = True
        elif any(kw in col_val for kw in ["expected", "answer", "reference", "उत्तर", "जवाब"]):
            expected_ans_col_idx = idx
            has_header = True

    start_row = 1 if has_header else 0

    parsed_questions = []
    warnings = []

    for row_num, row in enumerate(rows[start_row:], start=start_row + 1):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        q_text = ""
        exp_ans = ""

        if len(row) > question_col_idx and row[question_col_idx] is not None:
            q_text = str(row[question_col_idx]).strip()

        if len(row) > expected_ans_col_idx and row[expected_ans_col_idx] is not None:
            exp_ans = str(row[expected_ans_col_idx]).strip()

        if not q_text:
            warnings.append(f"Row {row_num}: Skipped because Question text is empty.")
            continue

        if not exp_ans:
            warnings.append(f"Row {row_num}: Expected answer empty — set to flexible default.")
            exp_ans = "Any Answer"

        q_id = f"q_{uuid.uuid4().hex[:8]}"
        parsed_questions.append({
            "question_id": q_id,
            "text": q_text,
            "expected_answer": exp_ans,
            "type": "text",
            "required": True,
            "options": [],
            "conditional_rules": [],
        })

    if not parsed_questions:
        raise HTTPException(
            status_code=400,
            detail="No valid questions could be extracted from the Excel file. Please ensure there is at least one non-empty Question row.",
        )

    created_doc = None
    if save_to_db:
        clean_title = title or (file.filename.rsplit(".", 1)[0].replace("_", " ").title()) or "Imported Script"
        new_q = Questionnaire(
            title=clean_title,
            description=f"Imported from {file.filename} ({len(parsed_questions)} questions)",
            agent_id=agent_id,
            tenant_id=tenant_id,
            version_id="v1",
            questions=parsed_questions,
            status="draft",
            is_published=False,
            version_history=[{
                "version_id": "v1",
                "published_at": datetime.now(UTC).isoformat(),
                "notes": f"Initial import from {file.filename}",
            }],
        )
        db.add(new_q)
        await db.flush()
        if agent_id:
            await sync_questionnaire_to_agent(db, agent_id, new_q)
        await db.commit()
        await db.refresh(new_q)
        created_doc = map_q_to_doc(new_q)

    return {
        "success": True,
        "filename": file.filename,
        "valid_count": len(parsed_questions),
        "questions": parsed_questions,
        "warnings": warnings,
        "questionnaire": created_doc,
    }


@router.get("/{questionnaire_id}")
async def get_questionnaire(questionnaire_id: str, db: AsyncSession = Depends(get_db)):
    stmt = select(Questionnaire).where(Questionnaire.id == questionnaire_id)
    res = await db.execute(stmt)
    q = res.scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="Questionnaire not found")
    return map_q_to_doc(q)


async def sync_questionnaire_to_agent(db: AsyncSession, agent_id: str, q: Questionnaire):
    """
    Connects a questionnaire to an agent: sets agent.questionnaire_id,
    injects formatted questions script into instructions & system prompt,
    and updates the voice workflow in Dograh.
    """
    if not agent_id:
        return
    stmt = select(Agent).where(Agent.id == agent_id)
    res = await db.execute(stmt)
    agent = res.scalar_one_or_none()
    if not agent:
        return

    agent.questionnaire_id = q.id
    q.agent_id = agent.id

    from app.services.agent_service import format_questionnaire_prompt_block
    import re
    if q.questions:
        q_block = format_questionnaire_prompt_block(q.title, q.questions)
        cleaned_instr = re.sub(r"\n\n### MANDATORY QUESTIONNAIRE SCRIPT[\s\S]*?(?=\n\n[A-Z#]|$)", "", agent.instructions or "").strip()
        agent.instructions = f"{cleaned_instr}\n\n{q_block}" if cleaned_instr else q_block
        agent.system_prompt = agent.instructions

    try:
        from app.services.dograh_service import dograh_service
        await dograh_service.sync_workflow_to_dograh(
            name=agent.name,
            purpose=agent.purpose,
            instructions=agent.instructions,
            allowed_topics=agent.allowed_topics or [],
            restricted_topics=agent.restricted_topics or [],
            behavior_rules=agent.behavior_rules or [],
            existing_workflow_id=agent.dograh_workflow_id,
            existing_trigger_path=agent.dograh_trigger_path,
            voice_id=agent.voice_id,
            speaking_speed=agent.speaking_speed or 1.0,
            language_mode=agent.language_mode or "auto_match",
            primary_language=agent.primary_language or "multi",
            voice_style=agent.voice_style or "Friendly + Professional + Conversational",
            tts_provider=agent.tts_provider or "dograh",
            company_name=agent.company_name,
            company_description=agent.company_description,
            greeting_message=agent.greeting_message,
        )
    except Exception as e:
        logger.warning(f"Error syncing agent workflow to Dograh: {e}")


@router.post("", status_code=201)
async def create_questionnaire(payload: QuestionnairePayload, db: AsyncSession = Depends(get_db)):
    resolved_questions = extract_questions_from_payload(payload)
    q = Questionnaire(
        title=payload.title or "New Questionnaire",
        description=payload.description,
        agent_id=payload.agent_id,
        tenant_id=payload.tenant_id,
        version_id=payload.version_id or "v1",
        questions=resolved_questions,
        status=payload.status or "draft",
        is_published=payload.is_published or False,
        version_history=[{"version_id": payload.version_id or "v1", "published_at": datetime.now(UTC).isoformat()}],
    )
    db.add(q)
    await db.flush()

    if payload.agent_id:
        await sync_questionnaire_to_agent(db, payload.agent_id, q)

    await db.commit()
    await db.refresh(q)
    return map_q_to_doc(q)


@router.put("/{questionnaire_id}")
async def update_questionnaire(questionnaire_id: str, payload: QuestionnairePayload, db: AsyncSession = Depends(get_db)):
    stmt = select(Questionnaire).where(Questionnaire.id == questionnaire_id)
    res = await db.execute(stmt)
    q = res.scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="Questionnaire not found")

    if payload.title is not None:
        q.title = payload.title
    if payload.description is not None:
        q.description = payload.description
    if payload.agent_id is not None:
        q.agent_id = payload.agent_id
    if payload.questions or payload.sections:
        q.questions = extract_questions_from_payload(payload)
    elif payload.questions is not None and not payload.sections:
        q.questions = payload.questions
    if payload.status is not None:
        q.status = payload.status
    if payload.is_published is not None:
        q.is_published = payload.is_published

    q.updated_at = datetime.now(UTC)

    target_agent_id = payload.agent_id if payload.agent_id is not None else q.agent_id
    if target_agent_id:
        await sync_questionnaire_to_agent(db, target_agent_id, q)

    await db.commit()
    await db.refresh(q)
    return map_q_to_doc(q)


@router.post("/{questionnaire_id}/publish")
async def publish_questionnaire(questionnaire_id: str, payload: Dict[str, Any] = None, db: AsyncSession = Depends(get_db)):
    stmt = select(Questionnaire).where(Questionnaire.id == questionnaire_id)
    res = await db.execute(stmt)
    q = res.scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="Questionnaire not found")

    q.is_published = True
    q.status = "published"
    notes = (payload or {}).get("notes", "Published update")
    history = list(q.version_history or [])
    history.append({
        "version_id": f"v{len(history) + 1}",
        "published_at": datetime.now(UTC).isoformat(),
        "notes": notes,
    })
    q.version_history = history
    q.updated_at = datetime.now(UTC)
    if q.agent_id:
        await sync_questionnaire_to_agent(db, q.agent_id, q)
    await db.commit()
    await db.refresh(q)
    return map_q_to_doc(q)


@router.post("/{questionnaire_id}/duplicate")
async def duplicate_questionnaire(questionnaire_id: str, db: AsyncSession = Depends(get_db)):
    stmt = select(Questionnaire).where(Questionnaire.id == questionnaire_id)
    res = await db.execute(stmt)
    q = res.scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="Questionnaire not found")

    new_q = Questionnaire(
        title=f"{q.title} (Copy)",
        description=q.description,
        agent_id=q.agent_id,
        tenant_id=q.tenant_id,
        version_id="v1",
        questions=q.questions or [],
        status="draft",
        is_published=False,
        version_history=[{"version_id": "v1", "published_at": datetime.now(UTC).isoformat(), "notes": "Duplicated copy"}],
    )
    db.add(new_q)
    await db.commit()
    await db.refresh(new_q)
    return map_q_to_doc(new_q)


@router.post("/assign")
async def assign_questionnaire(payload: Dict[str, Any], db: AsyncSession = Depends(get_db)):
    agent_id = payload.get("agent_id")
    questionnaire_id = payload.get("questionnaire_id")
    if not agent_id or not questionnaire_id:
        raise HTTPException(status_code=400, detail="agent_id and questionnaire_id required")

    q_stmt = select(Questionnaire).where(Questionnaire.id == questionnaire_id)
    q_res = await db.execute(q_stmt)
    q_obj = q_res.scalar_one_or_none()
    if not q_obj:
        raise HTTPException(status_code=404, detail="Questionnaire not found")

    await sync_questionnaire_to_agent(db, agent_id, q_obj)
    await db.commit()
    return {"success": True, "message": "Questionnaire assigned and connected to agent system prompt"}


@router.delete("/{questionnaire_id}")
async def delete_questionnaire(questionnaire_id: str, db: AsyncSession = Depends(get_db)):
    """
    Permanently delete a questionnaire script from the database.
    Unlinks it from any assigned agents first.
    """
    stmt = select(Questionnaire).where(Questionnaire.id == questionnaire_id)
    res = await db.execute(stmt)
    q = res.scalar_one_or_none()
    if not q:
        raise HTTPException(status_code=404, detail="Questionnaire not found")

    # 1. Unlink from any agents
    await db.execute(
        update(Agent)
        .where(Agent.questionnaire_id == questionnaire_id)
        .values(questionnaire_id=None)
    )

    # 2. Delete the questionnaire record permanently from PostgreSQL
    await db.delete(q)
    await db.commit()
    logger.info(f"Permanently deleted questionnaire {questionnaire_id} from PostgreSQL database")
    return {"success": True, "message": "Questionnaire deleted successfully", "id": questionnaire_id}


class BatchDeleteQuestionnairesRequest(BaseModel):
    ids: List[str]


@router.post("/batch-delete")
async def batch_delete_questionnaires(payload: BatchDeleteQuestionnairesRequest, db: AsyncSession = Depends(get_db)):
    """
    Permanently delete multiple questionnaires from the database in a single transaction.
    """
    if not payload.ids:
        return {"success": True, "deleted_count": 0}

    # 1. Unlink from any agents
    await db.execute(
        update(Agent)
        .where(Agent.questionnaire_id.in_(payload.ids))
        .values(questionnaire_id=None)
    )

    # 2. Delete all matching questionnaires
    stmt = select(Questionnaire).where(Questionnaire.id.in_(payload.ids))
    res = await db.execute(stmt)
    items = res.scalars().all()
    count = 0
    for q in items:
        await db.delete(q)
        count += 1

    await db.commit()
    logger.info(f"Permanently deleted {count} questionnaires in batch from PostgreSQL database")
    return {"success": True, "deleted_count": count}
